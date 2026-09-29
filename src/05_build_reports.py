"""
Step 5 - Recurring reporting pack (Excel, formula-driven) + preview charts for the README.
"""
from pathlib import Path
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
R = ROOT / "outputs" / "sql_results"
IMG = ROOT / "images"; IMG.mkdir(exist_ok=True)
REP = ROOT / "reports"; REP.mkdir(exist_ok=True)

trend = pd.read_csv(R / "q02_monthly_trend_mom.csv")
lob = pd.read_csv(R / "q03_pmpm_by_line_of_business.csv")
reasons = pd.read_csv(R / "q05_denial_reasons.csv")
sla = pd.read_csv(R / "q09_turnaround_sla_by_claim_type.csv")
dq = pd.read_csv(ROOT / "outputs" / "dq_test_results.csv")

# ------------------------------------------------------------------ Excel pack
NAVY, BLUE = "1F3864", "0000FF"
HDR = dict(font=Font(name="Arial", bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=NAVY))
BODY = Font(name="Arial", size=10)
thin = Border(bottom=Side(style="thin", color="BFBFBF"))


def table(ws, df, r0=1, c0=1, money=(), pct=()):
    for j, col in enumerate(df.columns):
        cell = ws.cell(r0, c0 + j, col.replace("_", " ").title())
        cell.font, cell.fill = HDR["font"], HDR["fill"]
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(row):
            c = ws.cell(r0 + i, c0 + j, None if pd.isna(v) else v)
            c.font, c.border = Font(name="Arial", size=10, color=BLUE) if isinstance(v, (int, float)) else BODY, thin
            if df.columns[j] in money: c.number_format = '$#,##0;($#,##0)'
            if df.columns[j] in pct: c.number_format = '0.0%'
    for j, col in enumerate(df.columns):
        ws.column_dimensions[get_column_letter(c0 + j)].width = max(14, len(col) + 4)
    return r0 + len(df)


wb = Workbook()
ws = wb.active; ws.title = "Summary"
ws["A1"] = "Claims Performance Report - Jan 2024 to Dec 2025"; ws["A1"].font = Font(name="Arial", bold=True, size=14, color=NAVY)
ws["A2"] = "Synthetic data for portfolio purposes. Blue = source data; black = formula."; ws["A2"].font = Font(name="Arial", italic=True, size=9)

# Monthly sheet: raw inputs + formulas for PMPM and MoM
m = wb.create_sheet("Monthly Trend")
base = trend[["year_month", "claims", "paid", "member_months"]]
last = table(m, base, money=("paid",))
m["E1"], m["F1"] = "Paid PMPM", "PMPM MoM %"
for c in ("E1", "F1"):
    m[c].font, m[c].fill = HDR["font"], HDR["fill"]
for r in range(2, last + 1):
    m[f"E{r}"] = f"=C{r}/D{r}"; m[f"E{r}"].number_format = '$#,##0.00'
    m[f"F{r}"] = "" if r == 2 else f"=E{r}/E{r-1}-1"; m[f"F{r}"].number_format = '0.0%'
    for c in "EF": m[f"{c}{r}"].font = BODY
m.column_dimensions["E"].width = m.column_dimensions["F"].width = 14
ch = LineChart(); ch.title = "Paid PMPM by Month"; ch.height, ch.width = 8, 18; ch.y_axis.title = "$ PMPM"
ch.add_data(Reference(m, min_col=5, min_row=1, max_row=last), titles_from_data=True)
ch.set_categories(Reference(m, min_col=1, min_row=2, max_row=last)); m.add_chart(ch, "H2")

l = wb.create_sheet("Line of Business")
table(l, lob, money=("paid_pmpm",))
d = wb.create_sheet("Denials")
dl = table(d, reasons[["denial_reason", "denied_claims", "allowed_value_at_risk"]], money=("allowed_value_at_risk",))
d["D1"] = "% of Denials"; d["D1"].font, d["D1"].fill = HDR["font"], HDR["fill"]
for r in range(2, dl + 1):
    d[f"D{r}"] = f"=B{r}/SUM($B$2:$B${dl})"; d[f"D{r}"].number_format = '0.0%'; d[f"D{r}"].font = BODY
d[f"A{dl+1}"] = "Total"; d[f"B{dl+1}"] = f"=SUM(B2:B{dl})"; d[f"C{dl+1}"] = f"=SUM(C2:C{dl})"
d[f"C{dl+1}"].number_format = '$#,##0'
for c in "ABC": d[f"{c}{dl+1}"].font = Font(name="Arial", bold=True)
bc = BarChart(); bc.type = "bar"; bc.title = "Denied Claims by Reason"; bc.height, bc.width = 8, 16; bc.legend = None
bc.add_data(Reference(d, min_col=2, min_row=1, max_row=dl), titles_from_data=True)
bc.set_categories(Reference(d, min_col=1, min_row=2, max_row=dl)); d.add_chart(bc, "F2")

s = wb.create_sheet("SLA")
table(s, sla)
q = wb.create_sheet("Data Quality")
table(q, dq[["stage", "test_case_id", "rule", "severity", "rows_failed", "status"]])
q.column_dimensions["C"].width = 55

# Summary KPIs as formulas referencing the detail sheets
kpis = [
    ("Total claims", f"=SUM('Monthly Trend'!B2:B{last})", '#,##0'),
    ("Total paid", f"=SUM('Monthly Trend'!C2:C{last})", '$#,##0'),
    ("Member months", f"=SUM('Monthly Trend'!D2:D{last})", '#,##0'),
    ("Paid PMPM (overall)", "=B6/B7", '$#,##0.00'),
    ("Denied claims", f"=Denials!B{dl+1}", '#,##0'),
    ("Denial rate", "=B9/B5", '0.0%'),
    ("Allowed $ at risk from denials", f"=Denials!C{dl+1}", '$#,##0'),
    ("DQ tests failing on RAW", '=COUNTIFS(\'Data Quality\'!A:A,"RAW",\'Data Quality\'!F:F,"FAIL")', '0'),
    ("DQ tests failing on CLEAN", '=COUNTIFS(\'Data Quality\'!A:A,"CLEAN",\'Data Quality\'!F:F,"FAIL")', '0'),
]
ws["A4"], ws["B4"] = "KPI", "Value"
for c in ("A4", "B4"): ws[c].font, ws[c].fill = HDR["font"], HDR["fill"]
for i, (k, f, fmt) in enumerate(kpis, 5):
    ws[f"A{i}"], ws[f"B{i}"] = k, f
    ws[f"A{i}"].font = ws[f"B{i}"].font = BODY; ws[f"B{i}"].number_format = fmt; ws[f"A{i}"].border = ws[f"B{i}"].border = thin
ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 18
ws["A15"] = "Note: Dec 2025 paid is understated - claims still pending / not yet received (claims run-out lag)."
ws["A15"].font = Font(name="Arial", italic=True, size=9)
wb.save(REP / "Monthly_Claims_Report.xlsx")

# ------------------------------------------------------------------ README charts
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.spines.top": False, "axes.spines.right": False})
C1, C2, GREY = "#2E5A88", "#D9822B", "#9AA5B1"

fig, ax = plt.subplots(figsize=(10, 3.8))
ax.plot(trend.year_month, trend.paid_pmpm, color=C1, lw=2, marker="o", ms=4)
ax.set_title("Paid PMPM by month - winter peaks every Dec-Feb", loc="left", fontweight="bold")
ax.set_ylabel("$ per member per month"); ax.tick_params(axis="x", rotation=60); ax.grid(axis="y", alpha=.3)
ax.annotate("Dec-25 incomplete\n(claims lag)", xy=(len(trend) - 1, trend.paid_pmpm.iloc[-1]), xytext=(-110, 30),
            textcoords="offset points", arrowprops=dict(arrowstyle="->", color=GREY), color="#555")
fig.tight_layout(); fig.savefig(IMG / "pmpm_trend.png", dpi=150); plt.close()

den = pd.read_csv(R / "q04_denials_by_network_specialty.csv")
piv = den.pivot(index="specialty", columns="network_status", values="denial_rate_pct").sort_values("Out-of-Network")
fig, ax = plt.subplots(figsize=(8, 4.5))
y = range(len(piv))
ax.barh([i + .2 for i in y], piv["Out-of-Network"], .4, color=C2, label="Out-of-Network")
ax.barh([i - .2 for i in y], piv["In-Network"], .4, color=C1, label="In-Network")
ax.set_yticks(list(y)); ax.set_yticklabels(piv.index); ax.set_xlabel("Denial rate %")
ax.set_title("Denial rate: out-of-network ~2-3x in-network", loc="left", fontweight="bold"); ax.legend(frameon=False)
fig.tight_layout(); fig.savefig(IMG / "denials_by_network.png", dpi=150); plt.close()

raw = dq[dq.stage == "RAW"].sort_values("rows_failed")
fig, ax = plt.subplots(figsize=(8, 4.8))
ax.barh(raw.test_case_id + "  " + raw.rule.str.slice(0, 42), raw.rows_failed,
        color=[C2 if s == "FAIL" else GREY for s in raw.status])
ax.set_title("Data-quality test results on the RAW extract (all 15 pass after cleansing)", loc="left", fontweight="bold", fontsize=10)
ax.set_xlabel("Failing rows")
fig.tight_layout(); fig.savefig(IMG / "dq_results.png", dpi=150); plt.close()
print("Excel pack + charts written")
